"""Project import from Excel (.xlsx), CSV (UTF-8, with or without BOM) or the JSON portfolio manifest.

plan() is a dry-run (writes nothing) that reports every row and every bad cell; apply() re-plans, backs up
the database, then creates (always UNPUBLISHED) and, only when asked, updates existing projects.
Used by the admin page and by tools/import_projects.py. See docs/CONTENT_GUIDE.md.
"""

import csv
import io
import json
from dataclasses import asdict, dataclass, field
from typing import get_args

from fastapi import HTTPException
from pydantic import ValidationError
from sqlalchemy import select
from sqlalchemy.orm import Session

from .bulk import backup_db
from .models import ImportLog, Project, User
from .schemas import Country, ProjectIn, ProjectType

# (column, required, help) in template order. "published" is deliberately absent: imports never publish.
COLUMNS: list[tuple[str, bool, str]] = [
    ("slug", False, "Web address part, e.g. villa-zamalek. Empty = made from the English title. Used to find existing projects."),
    ("title_en", True, "Project title in English."),
    ("title_ar", False, "Project title in Arabic."),
    ("country", True, "One of: " + ", ".join(get_args(Country))),
    ("type", True, "One of: " + ", ".join(get_args(ProjectType))),
    ("city_en", False, "City in English."),
    ("city_ar", False, "City in Arabic."),
    ("year", False, "Completion year, 4 digits."),
    ("area_m2", False, "Area in square metres (number)."),
    ("client_name", False, "Client name, only with the client's permission."),
    ("tagline_en", False, "One-line description (English)."),
    ("tagline_ar", False, "One-line description (Arabic)."),
    ("summary_en", False, "Short description (English)."),
    ("summary_ar", False, "Short description (Arabic)."),
    ("scope_en", False, "Scope of work (English)."),
    ("scope_ar", False, "Scope of work (Arabic)."),
    ("details_en", False, "Details (English)."),
    ("details_ar", False, "Details (Arabic)."),
    ("story_en", False, "Long story (English, Markdown)."),
    ("story_ar", False, "Long story (Arabic, Markdown)."),
    ("is_render", False, "yes = design visualization project, no = executed. Default no."),
    ("featured", False, "yes/no. Default yes."),
    ("sort_order", False, "Whole number; smaller numbers are listed first."),
]
NAMES = [c for c, _, _ in COLUMNS]
BOOL = {"is_render", "featured"}
MAX_ROWS = 2000


@dataclass
class Row:
    row: int  # spreadsheet row number (header = 1) or manifest index + 1
    slug: str = ""
    title_en: str = ""
    action: str = "error"  # create | update | skip | error
    errors: list[str] = field(default_factory=list)  # "column: problem"
    notes: list[str] = field(default_factory=list)
    values: dict = field(default_factory=dict, repr=False)


def _text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def read(filename: str, raw: bytes) -> list[tuple[int, dict]]:
    """Rows as (row number, {column: text}). Raises 400 for an unreadable file."""
    name = (filename or "").lower()
    if name.endswith(".json") or raw.lstrip()[:1] in (b"[", b"{"):
        try:
            data = json.loads(raw.decode("utf-8-sig"))
        except (UnicodeDecodeError, json.JSONDecodeError) as e:
            raise HTTPException(400, f"Not a valid JSON file: {e}") from None
        items = data.get("projects", []) if isinstance(data, dict) else data
        if not isinstance(items, list):
            raise HTTPException(400, "JSON must be a list of projects")
        return [(i + 1, {k: v for k, v in it.items()}) for i, it in enumerate(items) if isinstance(it, dict)]
    if raw[:4] == b"PK\x03\x04":
        from openpyxl import load_workbook

        try:
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        except Exception:  # noqa: BLE001 - any parse failure is the user's file, not a server error
            raise HTTPException(400, "Could not read this Excel file (save it as .xlsx)") from None
        ws = wb["Projects"] if "Projects" in wb.sheetnames else wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        header = [_text(h).lower() for h in next(it, [])]
        rows = []
        for n, values in enumerate(it, start=2):
            r = {h: values[i] for i, h in enumerate(header) if h and i < len(values)}
            if any(_text(v) for v in r.values()):
                rows.append((n, r))
        return rows
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise HTTPException(400, "The CSV is not UTF-8. In Excel use Save As → 'CSV UTF-8 (Comma delimited)'.") from None
    reader = csv.DictReader(io.StringIO(text))
    reader.fieldnames = [(h or "").strip().lower() for h in reader.fieldnames or []]
    return [(n, r) for n, r in enumerate(reader, start=2) if any((v or "").strip() for v in r.values() if isinstance(v, str))]


def _parse(raw: dict, row: Row) -> dict:
    """Known columns only, converted; problems go to row.errors as 'column: message'."""
    unknown = sorted(k for k in raw if k and k not in NAMES and k != "published")
    if unknown:
        row.notes.append("ignored columns: " + ", ".join(unknown[:12]))
    out: dict = {}
    for col in NAMES:
        if col not in raw:
            continue
        v = raw[col]
        if col in BOOL:
            if isinstance(v, bool):
                out[col] = v
                continue
            t = _text(v).lower()
            if t == "":
                continue
            if t in ("1", "yes", "y", "true"):
                out[col] = True
            elif t in ("0", "no", "n", "false"):
                out[col] = False
            else:
                row.errors.append(f"{col}: use yes or no")
            continue
        t = _text(v)
        if t == "":
            continue
        if col in ("year", "sort_order"):
            try:
                out[col] = int(float(t))
            except ValueError:
                row.errors.append(f"{col}: must be a whole number")
            continue
        if col == "area_m2":
            try:
                out[col] = float(t.replace(",", ""))
            except ValueError:
                row.errors.append(f"{col}: must be a number")
            continue
        out[col] = t.lower() if col in ("country", "type", "slug") else t
    return out


def plan(db: Session, rows: list[tuple[int, dict]], update_existing: bool = False) -> list[Row]:
    from .routers.admin import slugify

    if len(rows) > MAX_ROWS:
        raise HTTPException(413, f"At most {MAX_ROWS} rows per import")
    existing = {p.slug: p for p in db.scalars(select(Project))}
    seen: dict[str, int] = {}
    out: list[Row] = []
    for n, raw in rows:
        r = Row(row=n)
        out.append(r)
        values = _parse(raw, r)
        r.title_en = values.get("title_en", "")
        r.slug = slugify(values.get("slug") or r.title_en) if (values.get("slug") or r.title_en) else ""
        current = existing.get(r.slug)
        if current is None or not update_existing:
            try:  # full validation for a new project (required columns, allowed values, lengths)
                ProjectIn(**{**values, "published": False})
            except ValidationError as e:
                for err in e.errors():
                    col = str(err["loc"][0]) if err["loc"] else "row"
                    msg = "required" if err["type"] == "missing" else err["msg"]
                    r.errors.append(f"{col}: {msg}")
        else:  # partial update: only the filled cells are checked and applied
            try:
                base = {c: getattr(current, c) for c in NAMES if c != "slug"}
                ProjectIn(**{**base, **values, "published": False})
            except ValidationError as e:
                for err in e.errors():
                    r.errors.append(f"{err['loc'][0] if err['loc'] else 'row'}: {err['msg']}")
        if r.slug and r.slug in seen:
            r.errors.append(f"slug: same project as row {seen[r.slug]}")
        elif r.slug:
            seen[r.slug] = n
        if r.errors:
            continue
        r.values = values
        if current is None:
            r.action = "create"
        elif update_existing:
            r.action = "update"
            r.notes.append("only the filled cells are changed; published status is not changed")
        else:
            r.action, r.notes = "skip", [*r.notes, "already exists (tick 'update existing' to change it)"]
    return out


def apply(db: Session, rows: list[tuple[int, dict]], user: User, update_existing: bool = False,
          filename: str = "import") -> tuple[list[Row], str]:
    from .routers.admin import audit, unique_slug

    planned = plan(db, rows, update_existing)
    todo = [r for r in planned if r.action in ("create", "update")]
    backup = backup_db("pre-import") if todo else ""
    existing = {p.slug: p for p in db.scalars(select(Project))}
    for r in todo:
        vals = {k: v for k, v in r.values.items() if k != "slug"}
        if r.action == "create":
            p = Project(**ProjectIn(**{**vals, "published": False}).model_dump(exclude={"slug"}),
                        slug=unique_slug(db, r.slug))
            p.published = False  # imports never publish
            db.add(p)
            db.flush()
            audit(db, user, "import-create", "project", p.id)
        else:
            p = existing[r.slug]
            for k, v in vals.items():
                setattr(p, k, v)
            audit(db, user, "import-update", "project", p.id)
    db.add(ImportLog(user_email=user.email, kind="projects", filename=filename[:255],
                     created=sum(r.action == "create" for r in todo), updated=sum(r.action == "update" for r in todo),
                     skipped=sum(r.action == "skip" for r in planned), errors=sum(r.action == "error" for r in planned),
                     backup_file=backup))
    db.commit()
    return planned, backup


def summary(rows: list[Row]) -> dict:
    out = {"create": 0, "update": 0, "skip": 0, "error": 0}
    for r in rows:
        out[r.action] += 1
    return out


def rows_json(rows: list[Row]) -> list[dict]:
    return [{k: v for k, v in asdict(r).items() if k != "values"} for r in rows]


def template_xlsx() -> bytes:
    """Empty 'Projects' sheet with the columns, plus an 'Instructions' sheet (no sample data to import by mistake)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Projects"
    ws.append(NAMES)
    for c in ws[1]:
        c.font = Font(bold=True)
    for col, letter_list in (("country", get_args(Country)), ("type", get_args(ProjectType)), ("is_render", ("yes", "no")),
                             ("featured", ("yes", "no"))):
        idx = NAMES.index(col) + 1
        letter = ws.cell(row=1, column=idx).column_letter
        dv = DataValidation(type="list", formula1='"' + ",".join(letter_list) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{letter}2:{letter}{MAX_ROWS + 1}")
    ins = wb.create_sheet("Instructions")
    ins.append(["Column", "Required", "What to write"])
    for c in ins[1]:
        c.font = Font(bold=True)
    for name, req, help_ in COLUMNS:
        ins.append([name, "yes" if req else "", help_])
    ins.append([])
    for line in (
        "One project per row in the 'Projects' sheet. Leave a cell empty when the information is not confirmed: never guess.",
        "New projects are always created as drafts (not published). Publish them from the admin after review and approval.",
        "Existing projects (same slug) are skipped unless 'update existing' is ticked; then only the filled cells change.",
        "Upload the file in Admin → Import projects and run the dry run first: it lists every problem by row and column.",
        "Save as .xlsx, or as 'CSV UTF-8' so Arabic text is kept.",
    ):
        ins.append([line])
    ins.column_dimensions["A"].width = 14
    ins.column_dimensions["C"].width = 100
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
