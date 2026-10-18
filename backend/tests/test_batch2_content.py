"""Batch 2, phase B: photo quality + bulk upload, render/executed labels, focal point, approvals gate, readiness."""

import io
import os
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook
from PIL import Image

from app import bulk
from app.config import settings
from app.db import SessionLocal
from app.models import Approval, Project, ProjectImage
from app.models import Image as ImageRow

TOOLS = Path(__file__).resolve().parents[2] / "tools"


def img_bytes(w=1800, h=1200, color=(120, 90, 60), fmt="JPEG", exif=None) -> bytes:
    buf = io.BytesIO()
    kw = {"exif": exif} if exif is not None else {}
    Image.new("RGB", (w, h), color).save(buf, fmt, **kw)
    return buf.getvalue()


@pytest.fixture()
def standard(monkeypatch):
    """Real thresholds (the rest of the suite runs with relaxed ones)."""
    monkeypatch.setattr(settings, "image_min_long_edge", 1600)
    monkeypatch.setattr(settings, "image_min_cover_long_edge", 2400)


@pytest.fixture()
def gate(monkeypatch):
    monkeypatch.setattr(settings, "enforce_approvals", True)


def make_project(client, auth, slug, **kw):
    body = {"title_en": f"Test Project {slug}", "slug": slug, "country": "eg", "type": "office", "published": False} | kw
    r = client.post("/api/admin/projects", headers=auth, json=body)
    assert r.status_code == 201, r.text
    return r.json()


def counts():
    with SessionLocal() as db:
        return db.query(ProjectImage).count(), db.query(ImageRow).count(), sorted(p.name for p in settings.upload_dir.iterdir())


def approve(client, auth, subject_type, sid, status="approved"):
    a_type = "client_permission_to_publish" if subject_type == "project" else "logo_usage_permission"
    r = client.post("/api/admin/approvals", headers=auth, data={
        "subject_type": subject_type, "subject_id": str(sid), "approval_type": a_type, "status": status,
        "approved_by_name": "Test Approver", "approver_organization": "Example Co", "approval_date": "2026-10-01",
        "evidence_note": "e-mail (test)"})
    assert r.status_code == 201, r.text
    return r.json()


# ---------- naming / mapping ----------
def test_filename_convention_and_csv_override():
    m = bulk.NAME_RE.match("villa-test__03__executed__cover.JPG")
    assert m and m["slug"] == "villa-test" and m["nn"] == "03" and m["kind"].lower() == "executed" and m["cover"]
    assert bulk.NAME_RE.match("villa-test__03__photo.jpg") is None
    mapping = bulk.read_mapping("﻿filename,project_slug,order,kind,is_cover,alt_en,alt_ar\nDSC_1.jpg,villa-test,2,render,yes,Hall,ردهة\n".encode())
    assert mapping["dsc_1.jpg"]["project_slug"] == "villa-test" and mapping["dsc_1.jpg"]["alt_ar"] == "ردهة"


# ---------- bulk: dry-run, apply, duplicates, quality ----------
def test_bulk_dry_run_writes_nothing_then_apply_then_rerun_is_idempotent(client, auth, standard):
    make_project(client, auth, "bulk-a")
    files = [
        ("files", ("bulk-a__01__executed__cover.jpg", img_bytes(2600, 1700), "image/jpeg")),
        ("files", ("bulk-a__02__render.jpg", img_bytes(1800, 1200, (10, 10, 10)), "image/jpeg")),
        ("files", ("unknown-slug__01__render.jpg", img_bytes(1800, 1200, (1, 2, 3)), "image/jpeg")),
        ("files", ("bulk-a__03__render.jpg", img_bytes(900, 600, (5, 5, 5)), "image/jpeg")),  # too small
        ("files", ("bulk-a__04__executed.jpg", img_bytes(2600, 1700), "image/jpeg")),  # same bytes as 01
    ]
    before = counts()
    r = client.post("/api/admin/bulk-images/plan", headers=auth, files=files)
    assert r.status_code == 200, r.text
    plan = r.json()
    actions = {row["file"]: row["action"] for row in plan["rows"]}
    assert actions == {
        "bulk-a__01__executed__cover.jpg": "create", "bulk-a__02__render.jpg": "create",
        "unknown-slug__01__render.jpg": "reject", "bulk-a__03__render.jpg": "reject", "bulk-a__04__executed.jpg": "skip-duplicate",
    }
    assert next(x for x in plan["rows"] if x["file"].startswith("unknown"))["project"] == "NO MATCH"
    assert counts() == before  # dry-run: no rows, no files
    done = client.post("/api/admin/bulk-images/apply", headers=auth, data={"token": plan["token"]}).json()
    assert done["summary"] == {"create": 2, "skip-duplicate": 1, "reject": 2}
    assert done["backup"].endswith(".bak")
    detail = client.get("/api/admin/projects", headers=auth).json()
    p = next(x for x in detail if x["slug"] == "bulk-a")
    assert [i["kind"] for i in p["images"]] == ["executed", "render"] and p["images"][0]["is_cover"]
    # running the same batch again creates nothing
    again = client.post("/api/admin/bulk-images/plan", headers=auth, files=files[:2]).json()
    assert {row["action"] for row in again["rows"]} == {"skip-duplicate"}


def test_bulk_zip_and_csv_mapping(client, auth):
    make_project(client, auth, "bulk-zip")
    z = io.BytesIO()
    with zipfile.ZipFile(z, "w") as zf:
        zf.writestr("photos/DSC_0001.jpg", img_bytes(1800, 1200, (40, 50, 60)))
        zf.writestr("__MACOSX/._DSC_0001.jpg", b"junk")
    mapping = "filename,project_slug,order,kind,is_cover,alt_en,alt_ar\nDSC_0001.jpg,bulk-zip,1,render,1,Lobby,ردهة\n"
    r = client.post("/api/admin/bulk-images/plan", headers=auth, files=[
        ("files", ("batch.zip", z.getvalue(), "application/zip")),
        ("mapping", ("map.csv", mapping.encode("utf-8-sig"), "text/csv"))]).json()
    assert [(x["file"], x["action"], x["kind"], x["cover"]) for x in r["rows"]] == [("DSC_0001.jpg", "create", "render", True)]


def test_cli_dry_run_writes_nothing(client, auth, tmp_path):
    make_project(client, auth, "bulk-cli")
    (tmp_path / "bulk-cli__01__executed.jpg").write_bytes(img_bytes(1800, 1200, (90, 1, 1)))
    before = counts()
    env = os.environ | {"DATABASE_URL": settings.database_url, "UPLOAD_DIR": str(settings.upload_dir)}
    out = subprocess.run([sys.executable, str(TOOLS / "bulk_images.py"), "--dir", str(tmp_path)], env=env,
                         capture_output=True, text=True, check=True).stdout
    assert "create" in out and "DRY-RUN" in out
    assert counts() == before


def test_quality_rules_override_and_exif(client, auth, standard):
    p = make_project(client, auth, "quality")
    url = f"/api/admin/projects/{p['id']}/images"
    small = client.post(url, headers=auth, files={"files": ("s.jpg", img_bytes(1000, 700), "image/jpeg")})
    assert small.status_code == 422 and "too small" in small.json()["detail"]
    pano = client.post(url, headers=auth, files={"files": ("p.jpg", img_bytes(3000, 600), "image/jpeg")})
    assert pano.status_code == 422 and "panorama" in pano.json()["detail"]
    tall = client.post(url, headers=auth, files={"files": ("t.jpg", img_bytes(800, 2600), "image/jpeg")})
    assert tall.status_code == 422 and "too tall" in tall.json()["detail"]
    fake = client.post(url, headers=auth, files={"files": ("x.jpg", b"not an image at all", "image/jpeg")})
    assert fake.status_code == 400  # content sniffing, not the extension
    gps = Image.Exif()
    gps[0x8825] = {2: (30.0, 1.0, 1.0)}
    ok = client.post(url, headers=auth, data={"override_quality": "true", "kind": "executed"},
                     files={"files": ("g.jpg", img_bytes(1000, 700, exif=gps), "image/jpeg")})
    assert ok.status_code == 201, ok.text
    rec = ok.json()[0]
    assert rec["quality_override_by"] == "admin@test.com" and "too small" in rec["quality_warnings"]
    stored = settings.upload_dir / rec["image"]["src"].split("/")[2]
    for f in stored.iterdir():
        assert Image.open(f).getexif().get(0x8825) is None  # GPS stripped


def test_kind_is_required_to_publish_and_hidden_without_it(client, auth):
    p = make_project(client, auth, "kind-test")
    up = client.post(f"/api/admin/projects/{p['id']}/images", headers=auth, files={"files": ("a.jpg", img_bytes(), "image/jpeg")})
    iid = up.json()[0]["id"]
    assert up.json()[0]["kind"] is None
    body = {**p, "published": True}
    r = client.put(f"/api/admin/projects/{p['id']}", headers=auth, json=body)
    assert r.status_code == 422 and "Render or Executed" in r.json()["detail"]
    client.patch(f"/api/admin/project-images/{iid}", headers=auth, json={"kind": "render", "focal_x": 0.2, "focal_y": 0.8})
    assert client.put(f"/api/admin/projects/{p['id']}", headers=auth, json=body).status_code == 200
    d = client.get("/api/projects/kind-test").json()
    assert d["gallery"][0]["kind"] == "render" and d["gallery"][0]["is_render"] is True
    assert (d["cover"]["fx"], d["cover"]["fy"]) == (0.2, 0.8)
    assert "quality_override_by" not in d["gallery"][0] and "quality_warnings" not in d["gallery"][0]  # admin-only


# ---------- approvals gate ----------
def test_gate_hides_unapproved_and_history_is_append_only(client, auth, gate):
    p = make_project(client, auth, "gate-p", published=True)
    assert client.get("/api/projects/gate-p").status_code == 404
    assert "gate-p" not in {x["slug"] for x in client.get("/api/bootstrap").json()["projects"]}
    admin_view = next(x for x in client.get("/api/admin/projects", headers=auth).json() if x["slug"] == "gate-p")
    assert not admin_view["is_public"] and any("Approvals" in r for r in admin_view["block_reasons"])
    report = client.get("/api/admin/approvals/gate-report", headers=auth).json()
    assert "gate-p" in {x["slug"] for x in report["projects_hidden"]}

    approve(client, auth, "project", p["id"])
    assert client.get("/api/projects/gate-p").status_code == 200
    assert "gate-p" in client.get("/sitemap.xml").text
    approve(client, auth, "project", p["id"], status="revoked")
    assert client.get("/api/projects/gate-p").status_code == 404
    hist = client.get(f"/api/admin/approvals?subject_type=project&subject_id={p['id']}", headers=auth).json()
    assert [h["status"] for h in hist["history"]] == ["revoked", "approved"] and hist["current"] == "revoked"
    aid = hist["history"][-1]["id"]
    assert client.put(f"/api/admin/approvals/{aid}", headers=auth, json={}).status_code in (404, 405)
    assert client.delete(f"/api/admin/approvals/{aid}", headers=auth).status_code in (404, 405)
    with SessionLocal() as db:
        assert db.get(Approval, aid).status == "approved"  # the old row is untouched
    bad = client.post("/api/admin/approvals", headers=auth, data={"subject_type": "project", "subject_id": str(p["id"]),
                      "approval_type": "client_permission_to_publish", "status": "approved"})
    assert bad.status_code == 422  # who + how are required for an approval


def test_gate_newer_approval_of_another_project_does_not_hide_this_one(client, auth, gate):
    a = make_project(client, auth, "gate-a", published=True)
    b = make_project(client, auth, "gate-b", published=True)
    approve(client, auth, "project", a["id"])
    approve(client, auth, "project", b["id"])  # newer row, different subject
    assert client.get("/api/projects/gate-a").status_code == 200
    assert client.get("/api/projects/gate-b").status_code == 200
    slugs = {x["slug"] for x in client.get("/api/bootstrap").json()["projects"]}
    assert {"gate-a", "gate-b"} <= slugs


def test_deleted_project_approval_does_not_pass_to_a_reused_id(client, auth, gate):
    p = make_project(client, auth, "gate-del", published=True)
    approve(client, auth, "project", p["id"])
    assert client.delete(f"/api/admin/projects/{p['id']}", headers=auth).status_code == 204
    q = make_project(client, auth, "gate-new", published=True)
    with SessionLocal() as db:  # simulate SQLite handing out the same id again
        db.get(Project, q["id"]).id = p["id"]
        db.commit()
    assert client.get("/api/projects/gate-new").status_code == 404
    hist = client.get(f"/api/admin/approvals?subject_type=project&subject_id={p['id']}", headers=auth).json()
    assert [h["status"] for h in hist["history"]] == ["revoked", "approved"]


def test_gate_for_client_logos_and_evidence_file(client, auth, gate):
    from .conftest import make_png

    c = client.post("/api/admin/clients", headers=auth, json={"name": "Example Client Co", "published": True}).json()
    client.post(f"/api/admin/clients/{c['id']}/logo", headers=auth, files={"file": ("l.png", make_png(), "image/png")})
    names = lambda: {x["name"] for x in client.get("/api/bootstrap").json()["clients"]}  # noqa: E731
    assert "Example Client Co" not in names()
    r = client.post("/api/admin/approvals", headers=auth, data={
        "subject_type": "client", "subject_id": str(c["id"]), "approval_type": "logo_usage_permission", "status": "approved",
        "approved_by_name": "Test Person", "evidence_note": "letter (test)"},
        files={"evidence": ("permission.pdf", b"%PDF-1.4 test", "application/pdf")})
    assert r.status_code == 201 and r.json()["has_evidence"]
    assert "Example Client Co" in names()
    ev = client.get(f"/api/admin/approvals/{r.json()['id']}/evidence", headers=auth)
    assert ev.status_code == 200 and ev.headers["content-disposition"].startswith("attachment")
    assert client.get(f"/api/admin/approvals/{r.json()['id']}/evidence").status_code == 401


# ---------- readiness ----------
def test_readiness_is_computed_and_exports(client, auth):
    p = make_project(client, auth, "ready-1", title_ar="مشروع", summary_en="Text", summary_ar="نص")
    rows = client.get("/api/admin/readiness", headers=auth).json()
    row = next(r for r in rows["rows"] if r["slug"] == "ready-1")
    assert not row["ready"] and "cover_image" in row["missing_items"] and "approved" in row["missing_items"]
    up = client.post(f"/api/admin/projects/{p['id']}/images", headers=auth, data={"kind": "executed"},
                     files={"files": ("r.jpg", img_bytes(1800, 1200, (3, 3, 99)), "image/jpeg")}).json()
    client.patch(f"/api/admin/project-images/{up[0]['id']}", headers=auth, json={"alt_en": "=HYPERLINK(\"x\")", "alt_ar": "نص"})
    approve(client, auth, "project", p["id"])
    row = next(r for r in client.get("/api/admin/readiness", headers=auth).json()["rows"] if r["slug"] == "ready-1")
    assert row["ready"] and row["missing"] == 0
    x = client.get("/api/admin/readiness?format=xlsx", headers=auth)
    ws = load_workbook(io.BytesIO(x.content)).active
    assert [c.value for c in ws[1]][:3] == ["type", "name", "slug"]
    csv_text = client.get("/api/admin/readiness?format=csv", headers=auth).content.decode("utf-8-sig")
    assert "ready-1" in csv_text
