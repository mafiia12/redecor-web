"""Batch 2, phase C2: enquiries CRM (workflow, history, assignment, follow-ups, filters, export) and daily digest."""

import io
import os
import socket
import subprocess
import sys
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app import crm
from app.config import settings
from app.db import SessionLocal
from app.models import Enquiry, EnquiryEvent, JobRun, User

TOOLS = Path(__file__).resolve().parents[2] / "tools"


def make_enquiry(**kw) -> int:
    values = dict(name="Test Person", phone="+20 100 000 0000", email="test@example.com", country="eg", service="fitout",
                  message="Test message", lang="en", notify_status="sent") | kw
    with SessionLocal() as db:
        e = Enquiry(**values)
        db.add(e)
        db.commit()
        return e.id


def admin_id() -> int:
    with SessionLocal() as db:
        return db.query(User).filter_by(email="admin@test.com").one().id


def patch(client, auth, eid, **body):
    return client.patch(f"/api/admin/enquiries/{eid}", headers=auth, json=body)


def test_status_workflow_with_append_only_history(client, auth):
    eid = make_enquiry(name="Test Workflow")
    assert patch(client, auth, eid, status="won").status_code == 422  # new -> won skips the workflow
    for s in ("contacted", "quoted"):
        assert patch(client, auth, eid, status=s).status_code == 200
    r = patch(client, auth, eid, status="won", note="Signed (test)")
    assert r.status_code == 200 and r.json()["status"] == "won"
    events = r.json()["events"]
    assert [(e["from_value"], e["to_value"]) for e in events if e["kind"] == "status"] == [("quoted", "won"), ("contacted", "quoted"), ("new", "contacted")]
    assert events[0]["user_email"] == "admin@test.com" and events[0]["note"] == "Signed (test)" and events[0]["at"]
    assert patch(client, auth, eid, status="quoted").status_code == 422  # finished: only reopening to contacted
    assert patch(client, auth, eid, status="contacted").status_code == 200
    assert patch(client, auth, eid, status="closed").status_code == 422  # legacy value cannot be chosen
    # no endpoint edits or deletes history
    assert client.delete(f"/api/admin/enquiries/{eid}", headers=auth).status_code in (404, 405)
    with SessionLocal() as db:
        assert db.query(EnquiryEvent).filter_by(enquiry_id=eid).count() == 4  # 4 status changes (the note rides on one)


def test_legacy_closed_can_move_anywhere(client, auth):
    eid = make_enquiry(name="Test Legacy", status="closed")
    assert client.get(f"/api/admin/enquiries/{eid}", headers=auth).json()["status"] == "closed"
    assert patch(client, auth, eid, status="won").status_code == 200


def test_assign_follow_up_notes_and_overdue(client, auth):
    eid = make_enquiry(name="Test Assign")
    me = admin_id()
    r = patch(client, auth, eid, assigned_to_id=me, next_follow_up=str(crm.today() - timedelta(days=2)), note="Called, no answer")
    d = r.json()
    assert d["assigned_to"] == "admin@test.com" and d["overdue"] is True
    assert {e["kind"] for e in d["events"]} == {"assign", "follow_up", "note"}
    client.post("/api/admin/users", headers=auth, json={"email": "ed-crm@test.com", "role": "editor", "password": "Strong-pass-2026"})
    with SessionLocal() as db:
        ed = db.query(User).filter_by(email="ed-crm@test.com").one().id
    assert patch(client, auth, eid, assigned_to_id=ed).status_code == 422  # only admins can hold enquiries
    assert patch(client, auth, eid, assigned_to_id=None).json()["assigned_to"] == ""
    # a finished enquiry is never overdue
    patch(client, auth, eid, status="lost")
    assert client.get(f"/api/admin/enquiries/{eid}", headers=auth).json()["overdue"] is False


def test_filters_and_search(client, auth):
    a = make_enquiry(name="Test Filter Alpha", lang="ar", message="kitchen")
    b = make_enquiry(name="Test Filter Beta", lang="en", status="quoted",
                     created_at=datetime(2020, 1, 15, 10, tzinfo=timezone.utc))
    patch(client, auth, a, next_follow_up=str(crm.today() - timedelta(days=1)), assigned_to_id=admin_id())

    def ids(**params):
        r = client.get("/api/admin/enquiries", headers=auth, params=params)
        assert r.status_code == 200, r.text
        return {x["id"] for x in r.json()}

    assert a in ids(lang="ar") and b not in ids(lang="ar")
    assert b in ids(status="quoted") and a not in ids(status="quoted")
    assert a in ids(overdue="true") and b not in ids(overdue="true")
    assert a in ids(assignee="me") and b in ids(assignee="none") and a not in ids(assignee="none")
    assert ids(date_from="2020-01-15", date_to="2020-01-15") >= {b} and a not in ids(date_from="2020-01-15", date_to="2020-01-15")
    assert ids(q="kitchen") >= {a} and b not in ids(q="kitchen")
    assert a in ids(status="open")


def test_export_protects_against_formula_injection(client, auth):
    make_enquiry(name="=HYPERLINK(\"http://evil.example.com\")", message="+1 call me", email="@x.example.com")
    csv_body = client.get("/api/admin/enquiries?format=csv", headers=auth)
    assert csv_body.status_code == 200 and "attachment" in csv_body.headers["content-disposition"]
    text = csv_body.content.decode("utf-8-sig")
    assert "'=HYPERLINK" in text and ",=HYPERLINK" not in text and "'+1 call me" in text
    x = load_workbook(io.BytesIO(client.get("/api/admin/enquiries?format=xlsx", headers=auth).content))
    values = [c.value for row in x.active.iter_rows() for c in row]
    assert "'=HYPERLINK(\"http://evil.example.com\")" in values
    assert not any(isinstance(v, str) and v.startswith("=") for v in values)


def test_enquiries_are_admin_only(client, auth):
    client.post("/api/admin/users", headers=auth, json={"email": "ed-crm2@test.com", "role": "editor", "password": "Strong-pass-2026"})
    tok = client.post("/api/admin/login", json={"email": "ed-crm2@test.com", "password": "Strong-pass-2026"}).json()["token"]
    ed = {"Authorization": f"Bearer {tok}"}
    eid = make_enquiry(name="Test Private")
    for method, url in (("get", "/api/admin/enquiries"), ("get", f"/api/admin/enquiries/{eid}"),
                        ("get", "/api/admin/enquiries?format=csv"), ("get", "/api/admin/enquiries/digest-preview")):
        assert getattr(client, method)(url, headers=ed).status_code == 403
    assert client.patch(f"/api/admin/enquiries/{eid}", headers=ed, json={"status": "spam"}).status_code == 403


def test_source_page_and_language_are_stored(client):
    r = client.post("/api/enquiries", json=dict(name="Test Source", phone="+20 100 222 3333", email="s@example.com", country="eg",
                                                service="fitout", consent=True, lang="ar", source="/ar/projects/x", elapsed_ms=5000))
    assert r.status_code == 201
    with SessionLocal() as db:
        e = db.query(Enquiry).filter_by(name="Test Source").one()
        assert e.source_page == "/ar/projects/x" and e.lang == "ar"


# ---------- digest ----------
def _reset_job():
    with SessionLocal() as db:
        db.query(JobRun).delete()
        db.commit()


def test_digest_dry_run_lists_new_and_overdue_and_records_nothing(client, auth):
    _reset_job()
    yesterday_noon = datetime.combine(crm.today() - timedelta(days=1), datetime.min.time(), crm.tz()).replace(hour=12)
    make_enquiry(name="Test Digest New", created_at=yesterday_noon.astimezone(timezone.utc))
    old = make_enquiry(name="Test Digest Overdue", created_at=datetime(2021, 5, 1, tzinfo=timezone.utc))
    patch(client, auth, old, next_follow_up=str(crm.today() - timedelta(days=3)))
    text = crm.run_digest(dry_run=True)
    new_part, overdue_part = text.split("Overdue follow-ups")
    assert "Test Digest New" in new_part and "Test Digest Overdue" in overdue_part.split("Follow-ups due today")[0]
    with SessionLocal() as db:
        assert db.get(JobRun, crm.JOB) is None
    env = {**os.environ, "PYTHONPATH": str(TOOLS.parent / "backend")}
    r = subprocess.run([sys.executable, str(TOOLS / "digest.py"), "--dry-run"], capture_output=True, text=True, env=env)
    assert r.returncode == 0 and "Enquiries digest for" in r.stdout, r.stderr


def test_digest_without_smtp_logs_and_skips(monkeypatch, caplog):
    _reset_job()
    monkeypatch.setattr(settings, "smtp_host", "")
    assert crm.run_digest() == "skipped:smtp-not-configured"
    assert "SMTP is not configured" in caplog.text
    assert crm.run_digest() == "skipped:already-sent"  # logged once per day, not every minute


def test_digest_sent_once_per_day_without_files(client, monkeypatch):
    aiosmtpd = pytest.importorskip("aiosmtpd.controller")
    received = []

    class Handler:
        async def handle_DATA(self, server, session, envelope):  # noqa: N802
            received.append(envelope)
            return "250 OK"

    with socket.socket() as sock:
        sock.bind(("127.0.0.1", 0))
        port = sock.getsockname()[1]
    ctl = aiosmtpd.Controller(Handler(), hostname="127.0.0.1", port=port)
    ctl.start()
    try:
        for k, v in (("smtp_host", "127.0.0.1"), ("smtp_port", port), ("smtp_security", "none"), ("smtp_user", ""),
                     ("mail_from", "web@example.com"), ("notify_to", ["sales@example.com"])):
            monkeypatch.setattr(settings, k, v)
        _reset_job()
        assert crm.run_digest() == "sent"
        assert crm.run_digest() == "skipped:already-sent"
        assert len(received) == 1 and received[0].rcpt_tos == ["sales@example.com"]
        body = received[0].content
        assert b"Enquiries digest" in body and b"Content-Disposition: attachment" not in body
        assert crm.run_digest(force=True) == "sent" and len(received) == 2
    finally:
        ctl.stop()


def test_digest_schedule_uses_configured_time_zone(monkeypatch):
    monkeypatch.setattr(settings, "digest_tz", "Africa/Cairo")
    monkeypatch.setattr(settings, "digest_time", "08:30")
    cairo = crm.tz()
    assert not crm.digest_due(datetime(2026, 1, 10, 8, 29, tzinfo=cairo))
    assert crm.digest_due(datetime(2026, 1, 10, 8, 30, tzinfo=cairo))
    assert crm.digest_due(datetime(2026, 1, 10, 6, 31, tzinfo=timezone.utc))  # 08:31 in Cairo (UTC+2 in January)
    monkeypatch.setattr(settings, "digest_time", "")
    assert not crm.digest_due(datetime(2026, 1, 10, 23, 0, tzinfo=cairo))  # empty = off
    assert isinstance(crm.today(), date)
