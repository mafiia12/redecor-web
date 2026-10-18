#!/usr/bin/env python3
"""Cut the client logos out of the "Our Clients" pages of the portfolio PDFs and upload them.

  python tools/import_clients.py --pdf-dir "D:\\Master Web (Redecor)" --dry-run     # preview in ./clients_preview
  python tools/import_clients.py --pdf-dir "D:\\Master Web (Redecor)" --email ADMIN --replace-demo

Sources, in order of preference: the "OUR CLIENTS" slide of a PowerPoint file in --pdf-dir (original logo pictures, best
quality), then the "Our Clients" pages of the PDFs. White circle badges are removed so only the logo remains.

Logos come out as transparent PNGs. White/grey logos taken from a dark page are recoloured dark so they read on the
light client wall (the site shows them in grayscale anyway). To name them, put one client name per line in a text file
in the same order as the preview sheet and pass --names names.txt; otherwise they are called "Client 01", "Client 02"...
and you can rename them in /admin → Clients.
"""
import argparse
import getpass
import hashlib
import io
import os
import sys
from pathlib import Path

import numpy as np
import requests
from PIL import Image
from scipy import ndimage

try:
    import pymupdf as fitz
except ImportError:
    import fitz

# which pages hold the logos in each file (keyword in file name -> pages)
SOURCES = [("fitout", [5, 6]), ("breif", [5, 6]), ("brief", [5, 6])]
INK = (28, 28, 30)


def find_pdf(pdf_dir: Path, key: str):
    for f in sorted(pdf_dir.rglob("*.pdf")):
        if key in "".join(c for c in f.name.lower() if c.isalnum()):
            return f


def embedded_logos(doc, page):
    """Separate pictures on the page (skips page-sized backgrounds)."""
    out = []
    area = page.rect.width * page.rect.height
    for info in page.get_images(full=True):
        try:
            d = doc.extract_image(info[0])
            im = Image.open(io.BytesIO(d["image"])).convert("RGBA")
            if d.get("smask"):
                m = Image.open(io.BytesIO(doc.extract_image(d["smask"])["image"])).convert("L").resize(im.size)
                im.putalpha(m)
            if min(im.size) < 40 or im.width * im.height > 0.6 * area * 4:
                continue
            out.append(im)
        except Exception:
            continue
    return out


def segment_logos(page):
    """Render the page, treat everything unlike the background as ink, and cut out each blob."""
    zoom = 2400 / page.rect.width
    pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
    rgb = np.array(Image.open(io.BytesIO(pix.tobytes("png"))).convert("RGB")).astype(np.int16)
    h, w, _ = rgb.shape
    border = np.concatenate([rgb[0], rgb[-1], rgb[:, 0], rgb[:, -1]])
    bg = np.median(border, axis=0)
    dist = np.abs(rgb - bg).max(axis=2)
    mask = dist > 40
    grown = ndimage.binary_dilation(mask, iterations=max(6, w // 150))  # merge icon + wordmark
    labels, n = ndimage.label(grown)
    boxes = []
    for sl in ndimage.find_objects(labels):
        y0, y1, x0, x1 = sl[0].start, sl[0].stop, sl[1].start, sl[1].stop
        bw, bh = x1 - x0, y1 - y0
        if bw < w * 0.025 or bh < h * 0.02 or bw > w * 0.6:
            continue
        if y1 < h * 0.18:  # the "OUR CLIENTS" heading
            continue
        boxes.append((x0, y0, x1, y1))
    # reading order: group blobs into rows by their centre, then left to right
    boxes.sort(key=lambda b: (b[1] + b[3]) / 2)
    band = (np.median([b[3] - b[1] for b in boxes]) * 0.5) if boxes else 0
    rows, cur = [], []
    for b in boxes:
        if cur and (b[1] + b[3]) / 2 - (cur[-1][1] + cur[-1][3]) / 2 > band:
            rows.append(cur)
            cur = []
        cur.append(b)
    if cur:
        rows.append(cur)
    out = []
    for b in [b for r in rows for b in sorted(r, key=lambda q: q[0] + q[2])]:
        x0, y0, x1, y1 = b
        crop = rgb[y0:y1, x0:x1].astype(np.uint8)
        light = np.median(crop.reshape(-1, 3), axis=0).mean() > 200
        alpha = np.clip(dist[y0:y1, x0:x1] * 3, 0, 255).astype(np.uint8)
        im = Image.fromarray(np.dstack([crop, np.full_like(alpha, 255) if light else alpha]), "RGBA")
        out.append((im, bg))
    return out


def normalise(im: Image.Image, bg=None) -> Image.Image:
    im = knock_out_white(im)
    bbox = im.getchannel("A").point(lambda a: 255 if a > 20 else 0).getbbox()
    if bbox:
        im = im.crop(bbox)
    a = np.array(im).astype(np.float32)
    weight = a[..., 3:4] / 255.0
    if weight.sum() > 0:
        mean = (a[..., :3] * weight).sum(axis=(0, 1)) / weight.sum()
        sat = (mean.max() - mean.min()) / (mean.max() + 1e-6)
        lum = mean.mean()
        # pale, colourless logo (from a dark page): recolour so it shows on a light background
        if sat < 0.2 and lum > 190:
            a[..., 0], a[..., 1], a[..., 2] = INK
    im = Image.fromarray(a.astype(np.uint8), "RGBA")
    pad = max(4, im.width // 40)
    canvas = Image.new("RGBA", (im.width + 2 * pad, im.height + 2 * pad), (0, 0, 0, 0))
    canvas.paste(im, (pad, pad))
    if canvas.width > 600:
        canvas = canvas.resize((600, round(canvas.height * 600 / canvas.width)), Image.LANCZOS)
    return canvas


def fingerprint(im: Image.Image) -> np.ndarray:
    """Shape and colour signature (so two logos with the same silhouette but different colours stay separate)."""
    white = Image.new("RGBA", im.size, (255, 255, 255, 255))
    g = Image.alpha_composite(white, im.convert("RGBA")).convert("RGB").resize((16, 16), Image.BILINEAR)
    return np.array(g, dtype=np.float32).ravel() / 255.0



def pptx_logos(pdf_dir: Path):
    """Pictures on the slide titled "OUR CLIENTS", in reading order (rows top to bottom, then left to right).
    Handles plain pictures and shapes filled with a picture (round badges)."""
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError:
        print("(python-pptx not installed: skipping PowerPoint files)")
        return []

    def walk(shapes):
        for sh in shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from walk(sh.shapes)
            else:
                yield sh

    for f in sorted(pdf_dir.rglob("*.pptx")):
        if f.name.startswith("~$"):
            continue
        try:
            prs = Presentation(f)
        except Exception as e:
            print(f"(could not open {f.name}: {e})")
            continue
        for slide in prs.slides:
            shapes = list(walk(slide.shapes))
            text = " ".join(sh.text_frame.text for sh in shapes if sh.has_text_frame).upper()
            if "OUR CLIENTS" not in text:
                continue
            import hashlib as _h

            allp = []  # (shape, rId, blob hash)
            for sh in shapes:
                rids = sh._element.xpath(".//a:blip/@r:embed")
                if rids and min(sh.width, sh.height) > 150000 and sh.width < prs.slide_width * 0.5:
                    try:
                        blob = slide.part.related_part(rids[0]).blob
                    except Exception:
                        continue
                    allp.append((sh, rids[0], _h.md5(blob).hexdigest()))
            counts: dict[str, int] = {}
            for _, _, hx in allp:
                counts[hx] = counts.get(hx, 0) + 1
            # a picture repeated many times is the round badge/frame, not a logo
            badges = [c for c in allp if counts[c[2]] >= 3]
            cands = [c for c in allp if counts[c[2]] < 3]
            print(f"{f.name}: slide 'OUR CLIENTS' has {len(allp)} pictures ({len(badges)} repeated badge frames, {len(cands)} others)")
            if badges:  # keep only pictures that sit inside a badge (drops the Redecor logo and decorations)
                def inside(c):
                    cx, cy = c[0].left + c[0].width / 2, c[0].top + c[0].height / 2
                    return any(b[0].left <= cx <= b[0].left + b[0].width and b[0].top <= cy <= b[0].top + b[0].height for b in badges)
                cands = [c for c in cands if inside(c)]
            elif len(cands) > 6:  # no frames: skip anything in the left third (the Redecor logo)
                cands = [c for c in cands if c[0].left + c[0].width / 2 > prs.slide_width * 0.4]
            print(f"  {len(cands)} logo pictures")
            if len(cands) < 6:
                continue
            med_h = sorted(c[0].height for c in cands)[len(cands) // 2]
            band = med_h * 0.8
            cands.sort(key=lambda c: c[0].top + c[0].height / 2)
            rows, cur = [], [cands[0]]
            for c in cands[1:]:
                if (c[0].top + c[0].height / 2) - (cur[-1][0].top + cur[-1][0].height / 2) > band:
                    rows.append(cur)
                    cur = []
                cur.append(c)
            rows.append(cur)
            ordered = [c for r in rows for c in sorted(r, key=lambda q: q[0].left + q[0].width / 2)]
            print(f"  using {len(ordered)} logos in {len(rows)} rows")
            out, failed = [], 0
            for sh, rid, _ in ordered:
                try:
                    out.append(Image.open(io.BytesIO(slide.part.related_part(rid).blob)).convert("RGBA"))
                except Exception:
                    failed += 1  # e.g. vector (EMF/SVG) pictures that Pillow cannot read
            if failed:
                print(f"  ! {failed} pictures could not be read (vector format?)")
            return out
    print("(no PowerPoint slide with the client logos was found: using the PDF pages instead)")
    return []


def knock_out_white(im: Image.Image) -> Image.Image:
    """Logos placed on white (or white circle badges): make the white transparent and keep the logo colours."""
    a = np.array(im).astype(np.int16)
    border = np.concatenate([a[0], a[-1], a[:, 0], a[:, -1]])
    opaque = border[:, 3] > 200
    if opaque.sum() < len(border) * 0.5:
        return im  # already has a transparent background
    if (border[opaque][:, :3].min(axis=1) > 225).mean() < 0.85:
        return im  # not a white background
    dist = np.clip((255 - a[..., :3]).max(axis=2) * 3, 0, 255)
    alpha = np.minimum(dist, a[..., 3]).astype(np.uint8)
    h, w = alpha.shape
    yy, xx = np.ogrid[:h, :w]
    if 0.85 < w / h < 1.18:  # round badge: ignore anything outside the circle
        alpha[(xx - w / 2) ** 2 + (yy - h / 2) ** 2 > (min(w, h) / 2 * 0.96) ** 2] = 0
    return Image.fromarray(np.dstack([a[..., :3].astype(np.uint8), alpha]), "RGBA")


def png_bytes(im):
    buf = io.BytesIO()
    im.save(buf, "PNG", optimize=True)
    return buf.getvalue()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--publish", action="store_true", help="publish the logos immediately (only when every client approved)")
    ap.add_argument("--pdf-dir", required=True)
    ap.add_argument("--api", default="http://localhost:8000")
    ap.add_argument("--email")
    ap.add_argument("--password", default=os.getenv("ADMIN_PASSWORD"))
    ap.add_argument("--names", help="text file, one client name per line, same order as the preview")
    ap.add_argument("--max", type=int, default=40)
    ap.add_argument("--replace-demo", action="store_true", help="delete the [placeholder] clients first")
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()

    pdf_dir = Path(a.pdf_dir)
    logos: list[Image.Image] = []
    seen: list[np.ndarray] = []
    for im in pptx_logos(pdf_dir):
        im = normalise(im)
        fp = fingerprint(im)
        if not any(np.abs(fp - s).mean() < 0.015 for s in seen):
            seen.append(fp)
            logos.append(im)
    for key, pages in ([] if logos else SOURCES):
        path = find_pdf(pdf_dir, key)
        if not path:
            continue
        doc = fitz.open(path)
        for n in pages:
            if n > len(doc):
                continue
            page = doc[n - 1]
            items = [(im, None) for im in embedded_logos(doc, page)]
            mode = "embedded"
            if len(items) < 4:  # logos are one picture: cut them out of the rendered page
                items = segment_logos(page)
                mode = "segmented"
            print(f"{path.name} p.{n}: {len(items)} logos ({mode})")
            for im, bg in items:
                im = normalise(im, bg)
                fp = fingerprint(im)
                if any(np.abs(fp - s).mean() < 0.015 for s in seen):  # same logo already collected
                    continue
                seen.append(fp)
                logos.append(im)
    logos = logos[: a.max]
    print(f"\n{len(logos)} unique logos")
    if not logos:
        sys.exit("No logos found. Check --pdf-dir and the page numbers in SOURCES.")

    names = []
    names_file = Path(a.names) if a.names else Path(__file__).with_name("client_names.txt")
    if names_file.exists():
        names = [l.strip() for l in names_file.read_text(encoding="utf-8").splitlines() if l.strip()]
        if not a.names and len(names) != len(logos):  # bundled names only fit the PowerPoint slide (16 logos)
            names = []
    label = lambda i: names[i] if i < len(names) else f"Client {i + 1:02d}"

    if a.dry_run:
        d = Path("clients_preview")
        d.mkdir(exist_ok=True)
        cells = []
        for i, im in enumerate(logos):
            (d / f"{i + 1:02d}.png").write_bytes(png_bytes(im))
            cells.append(f'<figure><div><img src="{i + 1:02d}.png"></div><figcaption>{i + 1}. {label(i)}</figcaption></figure>')
        (d / "index.html").write_text(
            '<meta charset="utf-8"><title>Client logos</title><style>body{font:14px system-ui;margin:24px;background:#fafaf8}'
            '.g{display:grid;grid-template-columns:repeat(auto-fill,180px);gap:12px}figure{margin:0}'
            'figure div{height:96px;background:#fff;border:1px solid #ddd;border-radius:8px;display:grid;place-items:center}'
            'img{max-width:140px;max-height:56px;filter:grayscale(1);opacity:.8}figcaption{font-size:12px;color:#555;margin-top:4px}</style>'
            '<h1>Client logos (as shown on the site)</h1><div class="g">' + "".join(cells) + "</div>",
            encoding="utf-8",
        )
        print("Open clients_preview\\\\index.html to review. Write the names in a text file (one per line, same order) and pass --names.")
        return

    if not a.email:
        sys.exit("--email is required (admin account)")
    s = requests.Session()
    r = s.post(f"{a.api}/api/admin/login", json={"email": a.email, "password": a.password or getpass.getpass("Admin password: ")})
    if r.status_code != 200:
        sys.exit(f"Login failed: {r.text}")
    s.headers["Authorization"] = f"Bearer {r.json()['token']}"
    if a.replace_demo:
        for c in s.get(f"{a.api}/api/admin/clients").json():
            if c["name"].startswith("["):
                s.delete(f"{a.api}/api/admin/clients/{c['id']}")
    for i, im in enumerate(logos):
        # Imported logos stay hidden until each client's permission is confirmed and they are published in /admin.
        r = s.post(f"{a.api}/api/admin/clients", json={"name": label(i), "sort_order": i + 1, "published": a.publish})
        if r.status_code != 201:
            print("  ! create failed:", r.text)
            continue
        cid = r.json()["id"]
        r = s.post(f"{a.api}/api/admin/clients/{cid}/logo", files={"file": (f"{i + 1:02d}.png", png_bytes(im), "image/png")})
        print(f"  {'✓' if r.status_code == 200 else '!'} {label(i)}")
    if not a.publish:
        print("Imported as UNPUBLISHED. Publish each logo in /admin > Clients once the client has approved its use.")


if __name__ == "__main__":
    main()
